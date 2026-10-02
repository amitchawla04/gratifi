import re, collections
from merge import rows
from inventory import I, CATS
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import CellIsRule, FormulaRule

F = 'Arial'
H = Font(name=F, bold=True, color='FFFFFF', size=10); HF = PatternFill('solid', fgColor='17171A')
B = Font(name=F, size=10); BB = Font(name=F, size=10, bold=True); BLUE = Font(name=F, size=10, color='0000FF'); T = Font(name=F, size=16, bold=True)
WR = Alignment(wrap_text=True, vertical='top'); CEN = Alignment(horizontal='center', vertical='top')
thin = Border(bottom=Side(style='thin', color='E7E5E1'))
YEL = PatternFill('solid', fgColor='FFFF00')
MAXR = 6000
catname = dict(CATS); item = {i[0]: i for i in I}
ALIAS = {'RequestTracker': 'ClaimTracker', 'ApplicationTracker': 'ClaimTracker', 'MissingPointsClaim': 'ClaimTracker', 'MapView': 'Map', 'TerminalMap': 'Map',
         'PhotoUpload': 'Upload', 'EvidenceUpload': 'Upload', 'DocumentUpload': 'Upload', 'ReceiptCapture': 'Upload', 'MenuItemOptions': 'MenuCard',
         'PlanTable': 'InstalmentPlan', 'PaymentSchedule': 'InstalmentPlan', 'ServicePass': 'TicketWallet', 'QuietHours': 'AlertSettings', 'SellerRating': 'ReviewSummary'}
def norm_blocks(s):
    return re.sub(r'NEW:\s*([A-Za-z]+)', lambda m: 'NEW: ' + ALIAS.get(m.group(1), m.group(1)), s)
CORE_FL = {('Book a return flight', 'Dates fixed'), ('Book a return flight', 'Flexible dates'), ('Understand a fare', 'Unsure what is included'), ('Pay with points', 'Enough points'), ('Pay with points and card', 'Not enough points'), ('Pay by card only', 'Save points'), ('Payment declined', 'Card issue'), ('Price changed at pay', 'Final check'), ('Booking confirmed', 'Success'), ('Pick seats', 'Together'), ('Add bags', 'Checked bag'), ('Check in', 'Online'), ('Show boarding pass', 'At the airport'), ('Track the flight', 'Live'), ('Search fails', 'Supplier down'), ('Price changed', 'Between search and choose'), ('Sold out', 'Last seats'), ('Seat taken', 'Choosing'), ('Authentication fails', 'Face ID / code')}

wb = Workbook()
# ---------- Read me ----------
r = wb.active; r.title = 'Read me'
lines = [
 ('Gratifi: every use case', T),
 ('Everything a bank\'s customer can ask Gratifi to do, across every product, service and piece of content R360 and Gratifi provide. One row is one thing a customer wants done in one situation: what they say, what Gratifi does, which UI kit blocks it uses, what the systems decide, and what changes by market.', B),
 ('', B),
 ('What is covered', BB),
 ('19 categories and 105 products and services (sheet "Coverage"), built from: the Commerce in chat doc (18 categories and the R360 project list: Open Rewards, TripSure, PHANTOM, Airport VIP, iRewards 2.0, Piingle, new-age rewards, ZERO, Gratifi, R360 core), CACHE, and the Gratifi app sitemap (12 bank areas and every sub-area).', B),
 ('Nothing counts as covered until "Coverage" shows zero gaps. An item is covered when it has at least 6 use cases, at least 2 where nothing goes wrong and at least 2 where something does. Add a product to "Coverage" and it shows as a gap until its use cases exist.', B),
 ('', B),
 ('How a row is built', BB),
 ('Stage × Job × Constraint × Exception. Stages: Discover, Search, Choose, Pay, Add-ons, Before, During, Problem, After, Account. "Customer says" is one example phrasing; the agent is trained on many per row. "Rules" are decided by systems, never by the model.', B),
 ('"UI blocks" names components in the Gratifi UI kit. "NEW: Name" means the kit does not have that block yet; sheet "Kit gaps" lists every one and how often it is needed.', B),
 ('', B),
 ('Scoring (1 low to 5 high) and build waves', BB),
 ('Volume: how often customers ask. Value: revenue or loyalty impact. Risk: harm if wrong (money, rights, safety, regulation). Priority = Volume × Value ÷ Risk. Edit the blue cells to change the waves.', B),
]
for i, (t, f) in enumerate(lines, 1):
    c = r.cell(row=i, column=1, value=t); c.font = f; c.alignment = WR
base = len(lines) + 1
r.cell(row=base, column=1, value='Wave 1 needs priority of at least').font = B; r.cell(row=base, column=2, value=6)
r.cell(row=base + 1, column=1, value='Wave 1 needs risk no higher than').font = B; r.cell(row=base + 1, column=2, value=2)
r.cell(row=base + 2, column=1, value='Risk of at least this goes to wave 3 (legal and human sign-off first)').font = B; r.cell(row=base + 2, column=2, value=4)
for k in range(3): c = r.cell(row=base + k, column=2); c.font = BLUE; c.fill = YEL
W1P, W1R, W3R = f"'Read me'!$B${base}", f"'Read me'!$B${base + 1}", f"'Read me'!$B${base + 2}"
r.cell(row=base + 3, column=1, value='Core path (Y) rows are wave 1 even at high risk, because a customer cannot complete a booking without them; legal still signs off the high-risk ones before launch.').font = B
st = base + 5
r.cell(row=st, column=1, value='Status and cautions').font = BB
for k, n in enumerate([
 'Scores are a first pass for discussion. Re-score with real volumes from partner banks and PHANTOM and TripSure history.',
 'Compensation and consumer-law amounts are given only where set by law and known (UK261, EU261, UK Section 75). Anything else is named and marked "to confirm" for each bank\'s legal team. Market notes are prompts to check, not legal advice.',
 'Investments (C13) and insurance selling need licences per market; those rows are information-only until legal signs off.',
 'Arabic copy in the kit is a layout draft for native review. Suppliers, airlines and hotels in examples are fictional unless the inventory names the supply.',
 'To add a use case: add a row at the bottom of "Use cases", fill the item ID and scores; copy Priority and Wave down.'], 1):
    c = r.cell(row=st + k, column=1, value='• ' + n); c.font = B; c.alignment = WR
r.column_dimensions['A'].width = 120; r.column_dimensions['B'].width = 10

# ---------- Use cases ----------
ws = wb.create_sheet('Use cases')
heads = ['ID', 'Category', 'Item ID', 'Item', 'Stage', 'Job', 'Constraint', 'Exception', 'Customer says', 'Gratifi does', 'UI blocks (kit)', 'Rules: systems decide', 'Market notes', 'Volume', 'Value', 'Risk', 'Core path', 'Priority', 'Wave']
widths = [12, 20, 8, 24, 10, 24, 18, 18, 34, 40, 30, 34, 40, 8, 8, 8, 8, 9, 7]
for j, h in enumerate(heads, 1):
    c = ws.cell(row=1, column=j, value=h); c.font = H; c.fill = HF; c.alignment = Alignment(wrap_text=True, vertical='center')
    ws.column_dimensions[get_column_letter(j)].width = widths[j - 1]
stage_order = ['Discover', 'Search', 'Choose', 'Pay', 'Add-ons', 'Before', 'During', 'Problem', 'After', 'Account']
iorder = {i[0]: n for n, i in enumerate(I)}
rs = sorted(rows, key=lambda x: (iorder[x['item']], stage_order.index(x['stage']) if x['stage'] in stage_order else 99))
cnt = collections.Counter()
for n, x in enumerate(rs, 2):
    cnt[x['item']] += 1
    it = item[x['item']]
    core = 'Y' if it[1] == 'C01' and (x['job'], x['constraint']) in CORE_FL else ''
    vals = [f"{x['item']}-{cnt[x['item']]:03d}", catname[it[1]], x['item'], it[2], x['stage'], x['job'], x['constraint'], x['exception'], x['says'], x['does'], norm_blocks(x['blocks']), x['rules'], x['markets'], int(x['vol']), int(x['val']), int(x['risk']), core]
    for j, v in enumerate(vals, 1):
        c = ws.cell(row=n, column=j, value=v); c.font = BLUE if j in (14, 15, 16, 17) else B; c.alignment = CEN if j >= 14 else WR; c.border = thin
    ws.cell(row=n, column=18, value=f'=IF(OR(N{n}="",P{n}=""),"",ROUND(N{n}*O{n}/P{n},1))')
    ws.cell(row=n, column=19, value=f'=IF(R{n}="","",IF(Q{n}="Y",1,IF(P{n}>={W3R},3,IF(AND(R{n}>={W1P},P{n}<={W1R}),1,2))))')
    for j in (18, 19): c = ws.cell(row=n, column=j); c.font = BB if j == 19 else B; c.alignment = CEN; c.border = thin
last = len(rs) + 1
ws.freeze_panes = 'D2'; ws.auto_filter.ref = f'A1:S{last}'
dv = DataValidation(type='whole', operator='between', formula1=1, formula2=5, showErrorMessage=True, error='Use 1 to 5'); ws.add_data_validation(dv); dv.add(f'N2:P{MAXR}')
dvs = DataValidation(type='list', formula1='"' + ','.join(stage_order) + '"', showErrorMessage=True); ws.add_data_validation(dvs); dvs.add(f'E2:E{MAXR}')
ws.conditional_formatting.add(f'S2:S{MAXR}', CellIsRule(operator='equal', formula=['1'], fill=PatternFill('solid', fgColor='FFE7D9')))
ws.conditional_formatting.add(f'S2:S{MAXR}', CellIsRule(operator='equal', formula=['3'], fill=PatternFill('solid', fgColor='FCE6E1')))
U = lambda col: f"'Use cases'!${col}$2:${col}${MAXR}"

# ---------- Coverage ----------
cv = wb.create_sheet('Coverage', 1)
ch = ['Item ID', 'Category', 'Item', 'What it covers', 'Supply', 'Status as last recorded', 'Use cases', 'Nothing wrong', 'Something wrong', 'Pay rows', 'Handovers', 'Proactive', 'New blocks needed', 'Coverage']
for j, h in enumerate(ch, 1):
    c = cv.cell(row=1, column=j, value=h); c.font = H; c.fill = HF; c.alignment = Alignment(wrap_text=True, vertical='center')
for n, it in enumerate(I, 2):
    for j, v in enumerate([it[0], catname[it[1]], it[2], it[3], it[4], it[5]], 1):
        c = cv.cell(row=n, column=j, value=v); c.font = B; c.alignment = WR; c.border = thin
    cv.cell(row=n, column=7, value=f'=COUNTIF({U("C")},A{n})')
    cv.cell(row=n, column=8, value=f'=COUNTIFS({U("C")},A{n},{U("H")},"")')
    cv.cell(row=n, column=9, value=f'=G{n}-H{n}')
    cv.cell(row=n, column=10, value=f'=COUNTIFS({U("C")},A{n},{U("E")},"Pay")')
    cv.cell(row=n, column=11, value=f'=COUNTIFS({U("C")},A{n},{U("K")},"*Handoff*")')
    cv.cell(row=n, column=12, value=f'=COUNTIFS({U("C")},A{n},{U("I")},"(Gratifi starts it)*")')
    cv.cell(row=n, column=13, value=f'=COUNTIFS({U("C")},A{n},{U("K")},"*NEW:*")')
    cv.cell(row=n, column=14, value=f'=IF(AND(G{n}>=6,H{n}>=2,I{n}>=2),"Covered","GAP")')
    for j in range(7, 15): c = cv.cell(row=n, column=j); c.font = BB if j == 14 else B; c.alignment = CEN; c.border = thin
clast = len(I) + 1
cv.conditional_formatting.add(f'N2:N{clast + 50}', CellIsRule(operator='equal', formula=['"GAP"'], fill=PatternFill('solid', fgColor='FCE6E1'), font=Font(name=F, bold=True, color='C0301C')))
cv.conditional_formatting.add(f'N2:N{clast + 50}', CellIsRule(operator='equal', formula=['"Covered"'], fill=PatternFill('solid', fgColor='E2F3E9'), font=Font(name=F, bold=True, color='16794A')))
cv.freeze_panes = 'C2'; cv.auto_filter.ref = f'A1:N{clast}'
for j, w in enumerate((8, 22, 28, 34, 34, 20, 9, 9, 10, 8, 10, 9, 10, 10), 1): cv.column_dimensions[get_column_letter(j)].width = w
cv.cell(row=clast + 2, column=1, value='To add a product or service: add a row above this line with a new Item ID and copy the formulas in G to N. It shows GAP until its use cases exist.').font = B

# ---------- Summary ----------
sm = wb.create_sheet('Summary', 1)
sm['A1'] = 'Summary'; sm['A1'].font = T
sm['A3'] = 'Items with a coverage gap'; sm['A3'].font = BB
sm['B3'] = f"=COUNTIF(Coverage!$N$2:$N${clast + 50},\"GAP\")"; sm['B3'].font = Font(name=F, size=14, bold=True)
sm['A4'] = 'Use cases'; sm['B4'] = f'=COUNTA({U("C")})'
sm['A5'] = 'Products and services'; sm['B5'] = f'=COUNTIF(Coverage!$N$2:$N${clast + 50},"Covered")+COUNTIF(Coverage!$N$2:$N${clast + 50},"GAP")'
sm['A6'] = 'Kit blocks still to build'; sm['B6'] = f"=COUNTA('Kit gaps'!$A$2:$A$500)"
for a in ('A4', 'A5', 'A6', 'B4', 'B5', 'B6'): sm[a].font = B
sm.conditional_formatting.add('B3', CellIsRule(operator='greaterThan', formula=['0'], font=Font(name=F, size=14, bold=True, color='C0301C')))
sm.conditional_formatting.add('B3', CellIsRule(operator='equal', formula=['0'], font=Font(name=F, size=14, bold=True, color='16794A')))
hs = ['Category', 'Products', 'Use cases', 'Wave 1', 'Wave 2', 'Wave 3', 'Gaps', 'Average priority']
for j, h in enumerate(hs, 1):
    c = sm.cell(row=8, column=j, value=h); c.font = H; c.fill = HF
for k, (cid, cn) in enumerate(CATS, 9):
    sm.cell(row=k, column=1, value=cn).font = B
    sm.cell(row=k, column=2, value=f'=COUNTIF(Coverage!$B$2:$B${clast + 50},A{k})')
    sm.cell(row=k, column=3, value=f'=COUNTIF({U("B")},A{k})')
    for w in (1, 2, 3): sm.cell(row=k, column=3 + w, value=f'=COUNTIFS({U("B")},$A{k},{U("S")},{w})')
    sm.cell(row=k, column=7, value=f'=COUNTIFS(Coverage!$B$2:$B${clast + 50},A{k},Coverage!$N$2:$N${clast + 50},"GAP")')
    sm.cell(row=k, column=8, value=f'=IFERROR(ROUND(AVERAGEIF({U("B")},A{k},{U("R")}),1),0)')
    for j in range(2, 9): sm.cell(row=k, column=j).font = B
tr = 9 + len(CATS)
sm.cell(row=tr, column=1, value='Total').font = BB
for col in 'BCDEFG': sm[f'{col}{tr}'] = f'=SUM({col}9:{col}{tr - 1})'; sm[f'{col}{tr}'].font = BB
sm[f'H{tr}'] = f'=ROUND(AVERAGE({U("R")}),1)'; sm[f'H{tr}'].font = BB
sm.cell(row=tr + 2, column=1, value='Wave 1 is what the agent handles end to end first. Wave 3 needs legal sign-off and a person in the loop before launch.').font = B
for col, w in zip('ABCDEFGH', (32, 10, 10, 9, 9, 9, 8, 16)): sm.column_dimensions[col].width = w

# ---------- Kit gaps ----------
kg = wb.create_sheet('Kit gaps')
newc = collections.Counter(); newi = collections.defaultdict(set)
for x in rows:
    for nm in re.findall(r'NEW:\s*([A-Za-z]+)', norm_blocks(x['blocks'])): newc[nm] += 1; newi[nm].add(x['item'])
for j, h in enumerate(['Block to build', 'Use cases needing it', 'Items', 'Also called'], 1):
    c = kg.cell(row=1, column=j, value=h); c.font = H; c.fill = HF
inv = collections.defaultdict(list)
for a, b in ALIAS.items(): inv[b].append(a)
for n, (nm, k) in enumerate(newc.most_common(), 2):
    kg.cell(row=n, column=1, value=nm).font = BB
    kg.cell(row=n, column=2, value=f"=COUNTIF({U('K')},\"*NEW: {nm},*\")+COUNTIF({U('K')},\"*NEW: {nm}\")+COUNTIF({U('K')},\"*NEW: {nm} *\")").font = B
    kg.cell(row=n, column=3, value=', '.join(sorted(newi[nm]))).font = B
    kg.cell(row=n, column=4, value=', '.join(inv.get(nm, []))).font = B
    for j in (1, 2, 3, 4): kg.cell(row=n, column=j).alignment = WR
for col, w in zip('ABCD', (24, 12, 60, 36)): kg.column_dimensions[col].width = w
kg.freeze_panes = 'B2'

# ---------- Markets ----------
mk = wb.create_sheet('Markets')
mh = ['Market', 'Currency', 'Number and date format', 'How a payment is confirmed', 'Automatic card payment is called', 'Tax on fares', 'Flight disruption rule', 'Amounts set in law?', 'Week starts', 'Languages', 'To confirm with the bank']
mrows = [
 ('United Kingdom', 'GBP £', 'en-GB · 1,234.50 · Fri 16 Oct', 'Face ID in the bank app', 'Direct Debit', 'No VAT on flights; Air Passenger Duty in fares', 'UK261', 'Yes: £220, £350, £520 by distance', 'Monday', 'English', 'Point value; claims process'),
 ('Eurozone', 'EUR €', 'en-IE for English; local languages later', 'Face ID in the bank app (strong customer authentication)', 'direct debit', 'International flights usually VAT-exempt; local ticket taxes (to confirm)', 'EU261 (reform agreed June 2026, adoption pending)', 'Yes: €250, €400, €600 by distance', 'Monday', 'English first; German, French, Spanish, Italian, Dutch to add', 'Point value; which languages first'),
 ('India', 'INR ₹', 'en-IN · lakh grouping (1,00,000) · Fri 16 Oct', 'One-time code (other methods allowed under RBI authentication rules; to confirm per bank)', 'auto-debit', 'GST; GST invoice needs company GSTIN', 'DGCA CAR Section 3, Series M, Part IV', 'Yes, set by DGCA; legal team to confirm current amounts', 'Sunday', 'English first; Hindi next', 'Point value; authentication method; EMI terms'),
 ('UAE', 'AED', 'en-AE and ar-AE (right to left)', 'Face ID in the bank app', 'direct debit', 'VAT zero-rated on international flights (to confirm)', 'Airline policy (legal team to confirm any GCAA rule)', 'Legal team to confirm', 'Monday', 'English and Arabic (Arabic copy needs native sign-off)', 'Point value; Arabic copy'),
 ('Singapore', 'SGD S$', 'en-SG · 1,234.50', 'Face ID in the bank app', 'GIRO', 'GST zero-rated on international flights; airport levies', 'Airline policy; no statutory scheme', 'No', 'Sunday', 'English', 'Point value'),
 ('Malaysia', 'MYR RM', 'en-MY · 1,234.50', 'Approval in the bank\'s app', 'auto-debit', 'No SST on fares (to confirm); departure levy on international flights', 'Malaysian Aviation Consumer Protection Code 2016, enforced by CAAM', 'Legal team to confirm', 'Monday', 'English first; Malay next', 'Point value; app approval flow'),
]
for j, h in enumerate(mh, 1):
    c = mk.cell(row=1, column=j, value=h); c.font = H; c.fill = HF; c.alignment = Alignment(wrap_text=True, vertical='center')
for n, row in enumerate(mrows, 2):
    for j, v in enumerate(row, 1): c = mk.cell(row=n, column=j, value=v); c.font = B; c.alignment = WR; c.border = thin
for j, w in enumerate((16, 11, 26, 30, 18, 30, 30, 28, 10, 30, 28), 1): mk.column_dimensions[get_column_letter(j)].width = w
mk['A9'] = 'These settings live in the kit as data (Markets card). A new market is one more row here and one more entry in the kit; no redesign.'; mk['A9'].font = B

# ---------- Grammar ----------
g = wb.create_sheet('Grammar')
g['A1'] = 'The grammar'; g['A1'].font = T
stages = [('Discover', 'Before they know what they want'), ('Search', 'Looking for options'), ('Choose', 'Comparing and deciding'), ('Pay', 'Money or points move'), ('Add-ons', 'Extras on top of the main thing'), ('Before', 'Bought or booked, not yet used'), ('During', 'Using it now'), ('Problem', 'Something went wrong'), ('After', 'Done; refunds, reviews, repeats'), ('Account', 'The bank and points side')]
for j, h in enumerate(['Stage', 'Means'], 1): c = g.cell(row=3, column=j, value=h); c.font = H; c.fill = HF
for k, (a, b) in enumerate(stages, 4): g.cell(row=k, column=1, value=a).font = B; g.cell(row=k, column=2, value=b).font = B
g['D3'] = 'Most common exceptions'; g['D3'].font = H; g['D3'].fill = HF
exc = collections.Counter(x['exception'] for x in rows if x['exception'])
for k, (e, n) in enumerate(exc.most_common(40), 4): g.cell(row=k, column=4, value=e).font = B; g.cell(row=k, column=5, value=n).font = B
g['G3'] = 'Most common constraints'; g['G3'].font = H; g['G3'].fill = HF
con = collections.Counter(x['constraint'] for x in rows if x['constraint'])
for k, (e, n) in enumerate(con.most_common(40), 4): g.cell(row=k, column=7, value=e).font = B; g.cell(row=k, column=8, value=n).font = B
for col, w in zip('ABCDEFGH', (14, 34, 3, 30, 6, 3, 30, 6)): g.column_dimensions[col].width = w

wb.save('Gratifi-use-cases-all.xlsx')
print('rows', len(rs), 'items', len(I), 'new blocks', len(newc))
