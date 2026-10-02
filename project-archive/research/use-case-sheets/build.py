import sys; sys.path.insert(0, '.')
from cases import C
CORE = {('Book a return flight', 'Dates fixed'), ('Book a return flight', 'Flexible dates'), ('Understand a fare', 'Unsure what is included'), ('Pay with points', 'Enough points'), ('Pay with points and card', 'Not enough points'), ('Pay by card only', 'Save points'), ('Payment declined', 'Card issue'), ('Price changed at pay', 'Final check'), ('Booking confirmed', 'Success'), ('Pick seats', 'Together'), ('Add bags', 'Checked bag'), ('Check in', 'Online'), ('Show boarding pass', 'At the airport'), ('Track the flight', 'Live')}
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import CellIsRule

F = 'Arial'
H = Font(name=F, bold=True, color='FFFFFF', size=10)
HF = PatternFill('solid', fgColor='17171A')
B = Font(name=F, size=10); BB = Font(name=F, size=10, bold=True); BLUE = Font(name=F, size=10, color='0000FF'); T = Font(name=F, size=16, bold=True)
WR = Alignment(wrap_text=True, vertical='top'); CEN = Alignment(horizontal='center', vertical='top')
thin = Border(bottom=Side(style='thin', color='E7E5E1'))
wb = Workbook()

# ---------- Read me ----------
r = wb.active; r.title = 'Read me'
rows = [
 ('Gratifi: flight use cases', T),
 ('Every thing a customer might want done with a flight, written as one row each, with what Gratifi does, the UI blocks it uses from the Gratifi UI kit, what the systems decide, and what changes by market.', B),
 ('', B),
 ('How a row is built', BB),
 ('Stage × Job × Constraint × Exception. The same job ("Book a return flight") appears once per constraint that changes the answer, and once per exception that changes the flow.', B),
 ('Customer says is one real example; the agent is trained on many phrasings per row. Gratifi does is the answer shape. UI blocks name components in the kit. Rules are decided by systems, never by the model.', B),
 ('', B),
 ('Scoring (1 low to 5 high)', BB),
 ('Volume: how often customers ask. Value: revenue or loyalty impact. Risk: harm if wrong (money, rights, safety). Priority = Volume × Value ÷ Risk.', B),
 ('', B),
 ('Build waves (edit the blue cells to change them)', BB),
]
for i, (t, f) in enumerate(rows, 1):
    c = r.cell(row=i, column=1, value=t); c.font = f; c.alignment = WR
r['A12'] = 'Wave 1 needs priority of at least'; r['B12'] = 6
r['A13'] = 'Wave 1 needs risk no higher than'; r['B13'] = 2
r['A14'] = 'Rows with risk of at least this go to wave 3 (legal and human sign-off first)'; r['B14'] = 4
for a in ('B12', 'B13', 'B14'): r[a].font = BLUE; r[a].fill = PatternFill('solid', fgColor='FFFF00')
for a in ('A12', 'A13', 'A14'): r[a].font = B
r['A15'] = 'Core path (Y) rows are wave 1 whatever their score: a customer cannot book without them.'; r['A15'].font = B
r['A16'] = 'Status'; r['A16'].font = BB
notes = [
 'Scores are a first pass by Claude for discussion; the product team should re-score with real volumes from partner banks.',
 'Compensation amounts are named only where set by law (UK261, EU261). For India (DGCA), Malaysia (MAVCOM), Singapore and UAE the rule is named and the amount must come from each bank\'s legal team.',
 'Airlines in the kit are fictional. Market notes are prompts to check, not legal advice.',
 'To add a use case: add a row at the bottom of "Use cases" and copy the Priority and Wave formulas down.',
]
for i, n in enumerate(notes, 17):
    c = r.cell(row=i, column=1, value='• ' + n); c.font = B; c.alignment = WR
r.column_dimensions['A'].width = 110; r.column_dimensions['B'].width = 10

# ---------- Use cases ----------
ws = wb.create_sheet('Use cases')
heads = ['ID', 'Stage', 'Job', 'Constraint', 'Exception', 'Customer says', 'Gratifi does', 'UI blocks (kit)', 'Rules: systems decide', 'Market notes', 'Volume', 'Value', 'Risk', 'Core path', 'Priority', 'Wave']
widths = [8, 14, 24, 18, 16, 34, 40, 30, 34, 44, 8, 8, 8, 8, 9, 7]
for j, h in enumerate(heads, 1):
    c = ws.cell(row=1, column=j, value=h); c.font = H; c.fill = HF; c.alignment = Alignment(wrap_text=True, vertical='center')
    ws.column_dimensions[get_column_letter(j)].width = widths[j - 1]
ws.row_dimensions[1].height = 30
order = ['Inspire', 'Search', 'Choose', 'Pay', 'Extras', 'Before the trip', 'Day of travel', 'Disruption', 'After the trip', 'Points and card']
cs = sorted(C, key=lambda c: order.index(c['stage']))
pref = {'Inspire': 'INS', 'Search': 'SRC', 'Choose': 'CHO', 'Pay': 'PAY', 'Extras': 'EXT', 'Before the trip': 'PRE', 'Day of travel': 'DAY', 'Disruption': 'DIS', 'After the trip': 'AFT', 'Points and card': 'PTS'}
cnt = {}
for i, c in enumerate(cs, 2):
    cnt[c['stage']] = cnt.get(c['stage'], 0) + 1
    vals = [f"FL-{pref[c['stage']]}-{cnt[c['stage']]:02d}", c['stage'], c['job'], c['constraint'], c['exception'], c['says'], c['does'], c['blocks'], c['rules'], c['markets'], c['vol'], c['val'], c['risk'], 'Y' if (c['job'], c['constraint']) in CORE else '']
    for j, v in enumerate(vals, 1):
        cell = ws.cell(row=i, column=j, value=v); cell.font = BLUE if j in (11, 12, 13, 14) else B; cell.alignment = CEN if j >= 11 else WR; cell.border = thin
    ws.cell(row=i, column=15, value=f'=ROUND(K{i}*L{i}/M{i},1)').font = B
    ws.cell(row=i, column=16, value=f"=IF(N{i}=\"Y\",1,IF(M{i}>='Read me'!$B$14,3,IF(AND(O{i}>='Read me'!$B$12,M{i}<='Read me'!$B$13),1,2)))").font = BB
    for j in (15, 16): ws.cell(row=i, column=j).alignment = CEN; ws.cell(row=i, column=j).border = thin
last = len(cs) + 1
ws.freeze_panes = 'C2'; ws.auto_filter.ref = f'A1:P{last}'
dv = DataValidation(type='whole', operator='between', formula1=1, formula2=5, showErrorMessage=True, error='Use 1 to 5'); ws.add_data_validation(dv); dv.add(f'K2:M{last + 200}')
ws.conditional_formatting.add(f'P2:P{last}', CellIsRule(operator='equal', formula=['1'], fill=PatternFill('solid', fgColor='FFE7D9')))
ws.conditional_formatting.add(f'P2:P{last}', CellIsRule(operator='equal', formula=['3'], fill=PatternFill('solid', fgColor='FCE6E1')))

# ---------- Summary ----------
sm = wb.create_sheet('Summary', 1)
sm['A1'] = 'Summary'; sm['A1'].font = T
hs = ['Stage', 'Use cases', 'Wave 1', 'Wave 2', 'Wave 3', 'Average priority']
for j, h in enumerate(hs, 1):
    c = sm.cell(row=3, column=j, value=h); c.font = H; c.fill = HF
rng = lambda col: f"'Use cases'!${col}$2:${col}${last}"
for i, st in enumerate(order, 4):
    sm.cell(row=i, column=1, value=st).font = B
    sm.cell(row=i, column=2, value=f'=COUNTIF({rng("B")},A{i})').font = B
    for w in (1, 2, 3): sm.cell(row=i, column=2 + w, value=f'=COUNTIFS({rng("B")},$A{i},{rng("P")},{w})').font = B
    sm.cell(row=i, column=6, value=f'=IFERROR(ROUND(AVERAGEIF({rng("B")},A{i},{rng("O")}),1),0)').font = B
tr = 4 + len(order)
sm.cell(row=tr, column=1, value='Total').font = BB
for col in 'BCDE': sm[f'{col}{tr}'] = f'=SUM({col}4:{col}{tr - 1})'; sm[f'{col}{tr}'].font = BB
sm[f'F{tr}'] = f'=ROUND(AVERAGE({rng("O")}),1)'; sm[f'F{tr}'].font = BB
sm[f'A{tr + 2}'] = 'Wave 1 is what the agent handles end to end first. Wave 3 needs legal sign-off and a human in the loop before launch.'; sm[f'A{tr + 2}'].font = B
for col, w in zip('ABCDEF', (22, 12, 10, 10, 10, 16)): sm.column_dimensions[col].width = w

# ---------- Markets ----------
mk = wb.create_sheet('Markets')
mh = ['Market', 'Currency', 'Number and date format', 'How a payment is confirmed', 'Automatic card payment is called', 'Tax on fares', 'Flight disruption rule', 'Amounts set in law?', 'Week starts', 'Languages', 'To confirm with the bank']
mrows = [
 ('United Kingdom', 'GBP £', 'en-GB · 1,234.50 · Fri 16 Oct', 'Face ID in the bank app', 'Direct Debit', 'VAT; Air Passenger Duty in fares', 'UK261', 'Yes: £220, £350, £520 by distance', 'Monday', 'English', 'Point value; claims process'),
 ('Eurozone', 'EUR €', 'en-IE for English; local languages later', 'Face ID in the bank app (strong customer authentication)', 'direct debit', 'VAT; local ticket taxes', 'EU261', 'Yes: €250, €400, €600 by distance', 'Monday', 'English first; German, French, Spanish, Italian, Dutch to add', 'Point value; which languages first'),
 ('India', 'INR ₹', 'en-IN · lakh grouping (1,00,000) · Fri 16 Oct', 'One-time code sent by SMS', 'auto-debit', 'GST; GST invoice needs company GSTIN', 'DGCA passenger charter', 'Set by DGCA; legal team to confirm current amounts', 'Sunday', 'English first; Hindi next', 'Point value; OTP limits; EMI terms'),
 ('UAE', 'AED', 'en-AE and ar-AE (right to left)', 'Face ID in the bank app', 'direct debit', 'VAT', 'Airline policy (legal team to confirm any GCAA rule)', 'Legal team to confirm', 'Monday', 'English and Arabic (Arabic copy needs native sign-off)', 'Point value; Arabic copy'),
 ('Singapore', 'SGD S$', 'en-SG · 1,234.50', 'Face ID in the bank app', 'GIRO', 'GST', 'Airline policy; no statutory scheme', 'No', 'Sunday', 'English', 'Point value'),
 ('Malaysia', 'MYR RM', 'en-MY · 1,234.50', 'Approval in the bank\'s app', 'auto-debit', 'SST', 'MAVCOM consumer code', 'Legal team to confirm', 'Monday', 'English first; Malay next', 'Point value; app approval flow'),
]
for j, h in enumerate(mh, 1):
    c = mk.cell(row=1, column=j, value=h); c.font = H; c.fill = HF; c.alignment = Alignment(wrap_text=True, vertical='center')
for i, row in enumerate(mrows, 2):
    for j, v in enumerate(row, 1):
        c = mk.cell(row=i, column=j, value=v); c.font = B; c.alignment = WR; c.border = thin
for j, w in enumerate((16, 11, 26, 26, 18, 24, 24, 26, 10, 30, 28), 1): mk.column_dimensions[get_column_letter(j)].width = w
mk.row_dimensions[1].height = 32
mk['A9'] = 'These settings live in the kit as data (Markets card). A new market is one more row here and one more entry in the kit; no redesign.'; mk['A9'].font = B

# ---------- Grammar ----------
g = wb.create_sheet('Grammar')
g['A1'] = 'The grammar'; g['A1'].font = T
stages = [('Inspire', 'Before they know where'), ('Search', 'They know roughly what and when'), ('Choose', 'Comparing options and fares'), ('Pay', 'Money or points move'), ('Extras', 'Seats, bags, lounge, hotel, rides'), ('Before the trip', 'Booked, not yet travelling'), ('Day of travel', 'At or on the way to the airport'), ('Disruption', 'Something went wrong'), ('After the trip', 'Home again'), ('Points and card', 'The bank side of travel')]
cons = ['Dates fixed', 'Flexible dates', 'Budget in cash', 'Budget in points', 'Direct only', 'Time of day', 'Airline preference', 'Bag included', 'Travelling with children', 'Group of 5 or more', 'Premium cabin', 'Business trip', 'Accessibility needs', 'Pet', 'Separate tickets']
excs = ['No results', 'Supplier down', 'Price rose', 'Sold out', 'Payment declined', 'Authentication failed', 'Not enough points', 'Ticketing pending', 'Schedule change', 'Delay over 3 hours', 'Cancellation', 'Missed connection', 'Denied boarding', 'Diversion', 'Downgrade', 'Lost bag']
for j, h in enumerate(['Stage', 'Means', 'Constraints', 'Exceptions'], 1):
    c = g.cell(row=3, column=j, value=h); c.font = H; c.fill = HF
for i in range(max(len(stages), len(cons), len(excs))):
    if i < len(stages): g.cell(row=4 + i, column=1, value=stages[i][0]).font = B; g.cell(row=4 + i, column=2, value=stages[i][1]).font = B
    if i < len(cons): g.cell(row=4 + i, column=3, value=cons[i]).font = B
    if i < len(excs): g.cell(row=4 + i, column=4, value=excs[i]).font = B
for col, w in zip('ABCD', (18, 34, 26, 24)): g.column_dimensions[col].width = w
g['A22'] = 'Every new row should combine one stage, one job, at most one constraint and at most one exception. The full space is several hundred rows; this sheet holds the ones worth building first plus the high-risk ones that must be designed now.'
g['A22'].font = B; g['A22'].alignment = WR; g.merge_cells('A22:D24')

wb.save('Gratifi-flight-use-cases.xlsx')
print('saved', last - 1, 'rows')
